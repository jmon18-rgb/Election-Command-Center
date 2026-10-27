"""Official 2020 first-choice county history; UOCAVA stays unallocated.

No 2026 ballot-outstanding estimate or ranked-choice winner is derived here.
"""
import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_HASH = '0880031fa74d62dbbaf2e6a40380b7a3801033e66fcf8634b2e6f4b3e3fb96ce'
SOURCE_URL = 'https://mobile.maine.gov/sos/sites/maine.gov.sos/files/content/assets/ussenator1120.xlsx'
CENSUS_HASH = '3c337402b5c6e8d5aa26b4278ccf4edc8989f2683765b3ffbf22296cdb2df3a0'
HEADER = ['CTY', 'MUNICIPALITY', 'Collins, Susan Margaret', 'Gideon, Sara I.',
          'Linn, Max Patrick', 'Savage, Lisa', 'Others', 'BLANK', 'TBC']


def parse_rows(rows, reference):
    if list(rows[0][:9]) != HEADER or list(rows[2][2:6]) != [
            'Republican', 'Democratic', 'Independent', 'Independent']:
        raise ValueError('Official office/candidate/party header changed')
    sums = {code: [0] * 7 for code in reference}
    totals, seen, uocava, state = {}, set(), None, None
    for row in rows[3:]:
        code, town = row[:2]
        if not town:
            if any(v is not None for v in row[:9]):
                raise ValueError('Unlabeled data row')
            continue
        counts = list(row[2:9])
        if any(type(v) is not int or v < 0 for v in counts):
            raise ValueError('Invalid official count')
        if sum(counts[:6]) != counts[6]:
            raise ValueError('Candidates plus blanks do not equal ballot total')
        if code:
            if code not in reference or (code, town) in seen:
                raise ValueError('Unknown county or duplicate municipality')
            seen.add((code, town))
            sums[code] = [a + b for a, b in zip(sums[code], counts)]
        elif town.endswith(' Total') and town != 'State Total':
            key = town.removesuffix(' Total')
            if key not in reference or key in totals:
                raise ValueError('Unknown or duplicate county subtotal')
            totals[key] = counts
        elif town == 'STATE UOCAVA' and uocava is None:
            uocava = counts
        elif town == 'State Total' and state is None:
            state = counts
        else:
            raise ValueError('Unknown or duplicate statewide row')
    if set(totals) != set(reference) or totals != sums or uocava is None or state is None:
        raise ValueError('Incomplete or unreconciled county subtotals')
    if [sum(v[i] for v in sums.values()) + uocava[i] for i in range(7)] != state:
        raise ValueError('County plus unallocated UOCAVA does not equal state')

    def votes(values):
        return {'dem_votes': values[1], 'rep_votes': values[0],
                'other_votes': sum(values[2:5]), 'blank_ballots': values[5],
                'total_contest_votes': sum(values[:5])}

    counties = [{'county': reference[c]['county'], 'county_fips': reference[c]['fips'],
                 'records': {'total': votes(sums[c])}} for c in sorted(reference)]
    return {'schema': 'same-office-first-choice-county-history-v1',
            'race_code': 'ME-SEN', 'baseline_year': 2020, 'round': 1,
            'status': 'HISTORICAL_REFERENCE_NOT_2026_FORECAST', 'counties': counties,
            'statewide': votes(state), 'state_unallocated_uocava': votes(uocava),
            'uocava_policy': 'Retained separately; never assigned to counties or discarded.',
            'live_votes_accepted': False, 'may_move_national_needle': False,
            'limitations': ['Final first-choice history, not reporting chronology or calibration.',
                           'Contest votes are not an estimate of outstanding 2026 ballots.',
                           'Round-one counts must not be relabeled final RCV outcomes.']}


def build():
    original = ROOT / 'data/maine_historical_sources/2020-official-us-senate.xlsx'
    raw = original.read_bytes()
    census = (ROOT / 'data/texas_historical_sources/2024-census-counties.zip').read_bytes()
    if hashlib.sha256(raw).hexdigest() != SOURCE_HASH or hashlib.sha256(census).hexdigest() != CENSUS_HASH:
        raise ValueError('Pinned original source bytes changed')
    with zipfile.ZipFile(io.BytesIO(census)) as z:
        counties = [r for r in csv.DictReader(io.StringIO(z.read(
            '2024_Gaz_counties_national.txt').decode('utf-8-sig')), delimiter='\t')
                    if r['GEOID'].startswith('23')]
    reference = {r['NAME'][:3].upper(): {'county': r['NAME'].removesuffix(' County'),
                                      'fips': r['GEOID']} for r in counties}
    if len(reference) != 16:
        raise ValueError('Independent Census county reference incomplete')
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    if workbook.sheetnames != ['US Senator']:
        raise ValueError('Official workbook sheet changed')
    output = parse_rows(list(workbook.active.values), reference)
    if (output['statewide']['dem_votes'], output['statewide']['rep_votes']) != (347223, 417645):
        raise ValueError('Independent statewide total mismatch')
    output['source_url'], output['source_sha256'] = SOURCE_URL, SOURCE_HASH
    output['county_reference_sha256'] = CENSUS_HASH
    return output, reference


if __name__ == '__main__':
    output, reference = build()
    (ROOT / 'data/maine_senate_2020_county_baseline_official.json').write_text(
        json.dumps(output, indent=2) + '\n', encoding='utf-8')
    (ROOT / 'data/maine_county_fips_reference.json').write_text(json.dumps({
        'source': 'US Census 2024 national county gazetteer', 'source_sha256': CENSUS_HASH,
        'counties': [{'county': r['county'], 'fips': r['fips']} for r in reference.values()]
    }, indent=2) + '\n', encoding='utf-8')
    print('Maine 16 county subtotals plus separate statewide UOCAVA reconcile; no live input enabled')
